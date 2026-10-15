"""
core/parcela.py — modelo da parcela e navegação da campanha (RV-02 / RV-06).

- `Parcela`: experimento + X + Y + linha → identificador (`X3Y7L2`).
- `parse_identificador`: identificador → (x, y, linha), derivado do mesmo
  `config.PARCELA_FORMATO` (fica configurável).
- `carregar_plots`: lê a Plan1 da semeadora (.csv ou .xlsx) → ordem dos plots.
- `Navegador`: posição atual da campanha. "modo campanha" (segue a planilha)
  ou "modo manual" (incremento aritmético nas faixas de config). Persiste em
  disco para retomar onde parou.

Ref.: Rife & Poland (2014) "Field Book" — navegação por plots em ordem
definida no campo. docs/MELHORIAS.md frente #5.
"""

from __future__ import annotations

import csv
import json
import os
import re
import sys
from dataclasses import dataclass, asdict
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import config
from . import io_atomico


# ---------------------------------------------------------------------------
# identificador
# ---------------------------------------------------------------------------
def _regex_do_formato(formato):
    pat = re.escape(formato)
    pat = pat.replace(re.escape("{x}"), r"(?P<x>\d+)")
    pat = pat.replace(re.escape("{y}"), r"(?P<y>\d+)")
    pat = pat.replace(re.escape("{linha}"), r"(?P<linha>\d+)")
    return re.compile("^" + pat + "$")


def identificador(x, y, linha):
    return config.PARCELA_FORMATO.format(x=x, y=y, linha=linha)


def plot_id(x, y):
    return config.PLOT_FORMATO.format(x=x, y=y)


def parse_identificador(texto):
    """'X3Y7L2' -> (3, 7, 2). Retorna None se não casar com o formato."""
    m = _regex_do_formato(config.PARCELA_FORMATO).match(texto or "")
    if not m:
        return None
    return int(m.group("x")), int(m.group("y")), int(m.group("linha"))


@dataclass
class Parcela:
    experimento: str = ""
    x: int = 1
    y: int = 1
    linha: int = 1

    def identificador(self):
        return identificador(self.x, self.y, self.linha)

    def plot_id(self):
        return plot_id(self.x, self.y)


# ---------------------------------------------------------------------------
# planilha da semeadora
# ---------------------------------------------------------------------------
_COL_X = "Parcela Larg. (X)"
_COL_Y = "Parcela Compr.(Y)"
_COL_EXP = "ID Opcional"


def ler_planilha(caminho):
    """Lê a Plan1 (.csv ou .xlsx) como lista de dicts (cabeçalho -> valor),
    na ordem do arquivo. None se o arquivo não existe / não deu para ler."""
    if not caminho or not os.path.exists(caminho):
        return None
    try:
        if caminho.lower().endswith(".xlsx"):
            return list(_ler_xlsx(caminho))
        with open(caminho, newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    except Exception:
        return None


def carregar_plots(caminho):
    """Lê a Plan1 (.csv ou .xlsx). Retorna (lista_de_(x,y), experimento) na
    ordem do arquivo, ou None se o arquivo não existe / não deu para ler."""
    linhas = ler_planilha(caminho)
    if linhas is None:
        return None

    plots = []
    experimento = ""
    for d in linhas:
        try:
            x = int(float(d[_COL_X]))
            y = int(float(d[_COL_Y]))
        except (KeyError, TypeError, ValueError):
            continue
        plots.append((x, y))
        if not experimento and d.get(_COL_EXP):
            experimento = str(d[_COL_EXP]).strip()
    return (plots, experimento) if plots else None


def _ler_xlsx(caminho):
    import openpyxl  # só é necessário se a planilha for .xlsx
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb["Plan1"] if "Plan1" in wb.sheetnames else wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    cabecalho = [str(c).strip() if c is not None else "" for c in next(it)]
    for linha in it:
        yield dict(zip(cabecalho, linha))


# ---------------------------------------------------------------------------
# navegador
# ---------------------------------------------------------------------------
class Navegador:
    """Posição atual da campanha. Guarda experimento, plot corrente e linha."""

    def __init__(self, plots=None, experimento=""):
        self.plots = plots or None                 # [(x,y), ...] ou None
        self.experimento = experimento or config.EXPERIMENTO_PADRAO
        self.indice = 0                            # índice do plot (modo campanha)
        self.linha = 1
        self._x = config.PARCELA_X_MIN             # modo manual
        self._y = config.PARCELA_Y_MIN

    # --- construção ---
    @classmethod
    def do_config(cls):
        """Cria a partir da planilha de config (se existir) ou modo manual."""
        carregado = carregar_plots(config.PLANILHA_PROGRAMACAO)
        if carregado:
            plots, exp = carregado
            return cls(plots=plots, experimento=exp)
        return cls()

    # --- estado ---
    @property
    def modo(self):
        return "campanha" if self.plots else "manual"

    @property
    def total_plots(self):
        return len(self.plots) if self.plots else None

    @property
    def x(self):
        return self.plots[self.indice][0] if self.plots else self._x

    @property
    def y(self):
        return self.plots[self.indice][1] if self.plots else self._y

    def parcela(self):
        return Parcela(self.experimento, self.x, self.y, self.linha)

    def identificador(self):
        return self.parcela().identificador()

    def rotulo_plot(self):
        if self.plots:
            return f"plot {self.indice + 1} de {self.total_plots}"
        return f"plot {plot_id(self.x, self.y)}"

    # --- navegação ---
    def proxima(self):
        if self.linha < config.PARCELA_LINHAS:
            self.linha += 1
        else:
            self.linha = 1
            self._avancar_plot(+1)

    def anterior(self):
        if self.linha > 1:
            self.linha -= 1
        else:
            self.linha = config.PARCELA_LINHAS
            self._avancar_plot(-1)

    def _avancar_plot(self, passo):
        if self.plots:
            self.indice = max(0, min(self.total_plots - 1, self.indice + passo))
        else:
            # modo manual: Y anda primeiro, depois X (faixas de config)
            y = self._y + passo
            x = self._x
            if y > config.PARCELA_Y_MAX:
                y = config.PARCELA_Y_MIN
                x = min(config.PARCELA_X_MAX, x + 1)
            elif y < config.PARCELA_Y_MIN:
                y = config.PARCELA_Y_MAX
                x = max(config.PARCELA_X_MIN, x - 1)
            self._x, self._y = x, y

    def definir_manual(self, x=None, y=None, linha=None):
        """Modo manual: seta X/Y/linha direto (ignorado em modo campanha p/ X,Y)."""
        if linha is not None:
            self.linha = max(1, min(config.PARCELA_LINHAS, int(linha)))
        if not self.plots:
            if x is not None:
                self._x = max(config.PARCELA_X_MIN, min(config.PARCELA_X_MAX, int(x)))
            if y is not None:
                self._y = max(config.PARCELA_Y_MIN, min(config.PARCELA_Y_MAX, int(y)))

    # --- persistência (retoma onde parou) ---
    def to_dict(self):
        return {
            "experimento": self.experimento,
            "modo": self.modo,
            "indice": self.indice,
            "linha": self.linha,
            "x": self._x,
            "y": self._y,
            "salvo_em": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }

    def salvar(self, caminho=None):
        io_atomico.escrever_json(caminho or config.ARQUIVO_PARCELA, self.to_dict())

    def restaurar(self, caminho=None):
        """Recarrega posição de disco. Silencioso se não houver/for inválido."""
        caminho = caminho or config.ARQUIVO_PARCELA
        try:
            with open(caminho, encoding="utf-8") as f:
                d = json.load(f)
        except (FileNotFoundError, json.JSONDecodeError):
            return self
        if d.get("experimento"):
            self.experimento = d["experimento"]
        self.linha = int(d.get("linha", 1))
        if self.plots:
            self.indice = max(0, min(self.total_plots - 1, int(d.get("indice", 0))))
        else:
            self._x = int(d.get("x", config.PARCELA_X_MIN))
            self._y = int(d.get("y", config.PARCELA_Y_MIN))
        return self
