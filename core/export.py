"""
core/export.py — preenche as colunas `linha 1/2/3` da planilha da semeadora
com a contagem corrigida do rover (RV-06).

Fluxo:
1. lê o CSV canônico (`core/registro.py`);
2. soma `plantulas_corrigido` por (X, Y, linha) — opcionalmente filtrando
   por experimento;
3. copia a planilha de programação e escreve os totais nas colunas
   `linha 1` / `linha 2` / `linha 3` das linhas (X, Y) correspondentes.

Saída: `.xlsx` (preserva o resto da planilha) ou `.csv`.

Ref.: docs/MELHORIAS.md frentes #2 e #5; RV-06.
"""

from __future__ import annotations

import csv
import os
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import config
from core import registro
from core import parcela as parcela_mod

_COL_X = parcela_mod._COL_X          # "Parcela Larg. (X)"
_COL_Y = parcela_mod._COL_Y          # "Parcela Compr.(Y)"
_COL_EXP = parcela_mod._COL_EXP      # "ID Opcional"
_COLS_LINHA = {1: "linha 1", 2: "linha 2", 3: "linha 3"}


def agregar_por_linha(caminho_csv=None, experimento=None):
    """Retorna {(x, y): {1: total, 2: total, 3: total}} — soma de
    `plantulas_corrigido` por linha. `experimento` filtra (None = tudo)."""
    caminho_csv = caminho_csv or config.ARQUIVO_CSV
    acc = defaultdict(lambda: {1: 0, 2: 0, 3: 0})
    for d in registro.ler(caminho_csv):
        if experimento and (d.get("experimento") or "") != experimento:
            continue
        try:
            x = int(d["X"]); y = int(d["Y"]); linha = int(d["linha"])
            corr = int(d["plantulas_corrigido"])
        except (KeyError, TypeError, ValueError):
            continue
        if linha in (1, 2, 3):
            acc[(x, y)][linha] += corr
    return dict(acc)


def exportar(caminho_saida, caminho_csv=None, caminho_planilha=None, experimento=None):
    """Gera a planilha preenchida. Retorna (caminho_saida, n_plots_preenchidos)."""
    caminho_planilha = caminho_planilha or config.PLANILHA_PROGRAMACAO
    totais = agregar_por_linha(caminho_csv, experimento)

    if caminho_saida.lower().endswith(".xlsx"):
        n = _exportar_xlsx(caminho_planilha, caminho_saida, totais)
    else:
        n = _exportar_csv(caminho_planilha, caminho_saida, totais)
    return caminho_saida, n


def _valor(totais, x, y, linha):
    plot = totais.get((x, y))
    if not plot:
        return None
    v = plot.get(linha, 0)
    return v if v else None


def _exportar_csv(caminho_planilha, caminho_saida, totais):
    if not caminho_planilha or not os.path.exists(caminho_planilha):
        raise FileNotFoundError(f"planilha de programação não encontrada: {caminho_planilha}")
    if caminho_planilha.lower().endswith(".xlsx"):
        linhas = list(parcela_mod._ler_xlsx(caminho_planilha))
        campos = list(linhas[0].keys()) if linhas else []
    else:
        with open(caminho_planilha, newline="", encoding="utf-8-sig") as f:
            r = csv.DictReader(f)
            campos = r.fieldnames or []
            linhas = list(r)

    n = 0
    for d in linhas:
        try:
            x = int(float(d[_COL_X])); y = int(float(d[_COL_Y]))
        except (KeyError, TypeError, ValueError):
            continue
        preencheu = False
        for linha, col in _COLS_LINHA.items():
            v = _valor(totais, x, y, linha)
            if v is not None:
                d[col] = v
                preencheu = True
        n += preencheu

    os.makedirs(os.path.dirname(caminho_saida) or ".", exist_ok=True)
    with open(caminho_saida, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=campos)
        w.writeheader()
        w.writerows(linhas)
    return n


def _exportar_xlsx(caminho_planilha, caminho_saida, totais):
    import openpyxl
    if not caminho_planilha or not os.path.exists(caminho_planilha):
        raise FileNotFoundError(f"planilha de programação não encontrada: {caminho_planilha}")
    wb = openpyxl.load_workbook(caminho_planilha)
    ws = wb["Plan1"] if "Plan1" in wb.sheetnames else wb[wb.sheetnames[0]]

    cab = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    try:
        ix, iy = cab.index(_COL_X), cab.index(_COL_Y)
        icols = {n: cab.index(col) for n, col in _COLS_LINHA.items()}
    except ValueError as e:
        raise ValueError(f"coluna esperada não encontrada na Plan1: {e}")

    n = 0
    for row in ws.iter_rows(min_row=2):
        try:
            x = int(float(row[ix].value)); y = int(float(row[iy].value))
        except (TypeError, ValueError):
            continue
        preencheu = False
        for linha, icol in icols.items():
            v = _valor(totais, x, y, linha)
            if v is not None:
                row[icol].value = v
                preencheu = True
        n += preencheu

    os.makedirs(os.path.dirname(caminho_saida) or ".", exist_ok=True)
    wb.save(caminho_saida)
    return n
